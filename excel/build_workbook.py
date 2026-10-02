"""Builds mortgage_risk_model.xlsx: live scorecard, deciles, PSI, calibration and scenario formulas over 2,000 OOT loans.

Typed numbers live on Inputs and on the imported Python value sheets (Scorecard, Loan_Sample attributes,
Python_Ref, Curves_Python). Change the scenario selector on Inputs and every formula sheet recalculates.
Run: py -3 build_workbook.py
"""
import json
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "python"))
from mortgage_risk import config, reconcile_excel as rx  # noqa: E402
from mortgage_risk.columns import (AGE, CUR_UPB, FEATURES_PD, INCENTIVE, LOAN_ID, MTM_LTV, SNAPSHOT_DATE,  # noqa: E402
                                   TARGET_PD12, UNEMP, UNEMP_CHG_12M)

# ===== CONFIG (user inputs) =====
OUT_FILE = Path(__file__).resolve().parent / "mortgage_risk_model.xlsx"
AUTHOR = "Deepak Chaudhary"
HDR, FIRST = 4, 5                      # Loan_Sample header row and first data row
DEC_HDR = 9                            # Deciles header row
SCEN_HDR = 24                          # Scenario loan-table header row
CALC_FIRST = 6                         # Score_Calc first variable row
PY_PD, PY_POINTS, PY_GBM = "py_scorecard_pd", "py_scorecard_points", "py_gbm_pd"
SAMPLE_PY_PD, SAMPLE_PY_POINTS, SAMPLE_GBM = "scorecard_pd", "scorecard_points", "gbm_pd"
DEC_COLS = ("decile", "n", "events", "event_rate", "mean_score", "cum_events_share", "lift", "ks_at_decile")
SCEN12_COLS = ("default_12m_pct_upb", "prepay_12m_pct_upb", "avg_cpr_12m", "avg_cdr_12m")
# ===== END CONFIG =====

BOLD = Font(bold=True)
_CACHE = {}
SCEN_SHOCK_VARS = (UNEMP, UNEMP_CHG_12M, MTM_LTV, INCENTIVE)


def build(out_file=OUT_FILE, selector=1, write_meta=True):
    if "ctx" not in _CACHE:
        table, sample, coefs, psi_bins = rx.load_tables()
        p = rx.load_model_params()
        _CACHE["ctx"] = (table, sample, coefs, psi_bins, p, rx.reference_all(table, sample, coefs, psi_bins, p))
    table, sample, coefs, psi_bins, p, ref = _CACHE["ctx"]
    bins = ref["bins"]
    variables = list(bins)
    n = len(sample)
    last = FIRST + n - 1

    wb = openpyxl.Workbook()
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    addr = {}

    def reg(name, sheet, ref_):
        addr[name] = f"{sheet}!{ref_}"

    def head(ws, row, labels, col0=1):
        for j, t in enumerate(labels):
            ws.cell(row, col0 + j, t).font = BOLD

    ws_readme = wb.active
    ws_readme.title = "README"
    ws_in = wb.create_sheet("Inputs")
    ws_sc = wb.create_sheet("Scorecard")
    ws_calc = wb.create_sheet("Score_Calc")
    ws_ls = wb.create_sheet("Loan_Sample")
    ws_dec = wb.create_sheet("Deciles")
    ws_psi = wb.create_sheet("PSI")
    ws_cal = wb.create_sheet("Calibration")
    ws_scn = wb.create_sheet("Scenario")
    ws_chk = wb.create_sheet("Checks")
    ws_con = wb.create_sheet("Conclusions")
    ws_ref = wb.create_sheet("Python_Ref")
    ws_cur = wb.create_sheet("Curves_Python")

    # ============================================================ Inputs
    NAMES = {}
    row = [3]

    def put(key, label, value, note=""):
        r = row[0]
        ws_in.cell(r, 1, label)
        ws_in.cell(r, 2, value)
        ws_in.cell(r, 3, note)
        NAMES[key] = f"Inputs!$B${r}"
        row[0] += 1
        return r

    ws_in["A1"] = "Inputs (all typed numbers live here; synthetic data, not real loans)"
    ws_in["A1"].font = BOLD
    put("sel", "Scenario selector (1 to 6)", selector, "Matches the scenario table below; change this one cell")
    put("lgd", "LGD assumption", p["lgd"], "config.LGD_ASSUMPTION")
    put("pdo", "Scorecard PDO (points to double the odds)", config.SCORECARD_PDO, "config.SCORECARD_PDO")
    put("base_score", "Scorecard base score", config.SCORECARD_BASE_SCORE, "config.SCORECARD_BASE_SCORE")
    put("base_odds", "Scorecard base odds (good:bad)", config.SCORECARD_BASE_ODDS, "config.SCORECARD_BASE_ODDS")
    put("factor", "Scorecard factor", f"={NAMES['pdo']}/LN(2)", "PDO / ln 2 (formula)")
    put("offset", "Scorecard offset", f"={NAMES['base_score']}-{NAMES['factor']}*LN({NAMES['base_odds']})",
        "base score - factor x ln(base odds) (formula)")
    put("intercept", "Scorecard logit intercept", p["intercept"], "Python scorecard fit")
    put("pd_a", "PD Platt calibrator slope", p["pd_a"], "Python calibrator, applied to the scorecard margin")
    put("pd_b", "PD Platt calibrator intercept", p["pd_b"], "Python calibrator")
    put("clip", "Probability clip", p["prob_clip"], "models.PROB_CLIP")
    put("burnout", "Prepay BURNOUT assumption (months)", p["burnout"],
        "APPROXIMATION: no burnout history in the sample; one fixed value solved so base 12m prepay matches the Python GBM base")
    put("pp_a", "Prepay Platt calibrator slope", p["pp_a"], "Python calibrator on the prepay hinge logit")
    put("pp_b", "Prepay Platt calibrator intercept", p["pp_b"], "Python calibrator")
    put("season", "Seasoning ramp months", p["season_months"], "SEASONING_RAMP = min(age / this, 1)")
    put("mtm_cap", "Prepay MTM_LTV cap", p["mtm_cap"], "models.MTM_LTV_CAP")
    put("psi_floor", "PSI share floor", p["psi_floor"], "evaluation.PSI_FLOOR")
    put("psi_lo", "PSI band: stable below", config.PSI_BANDS[0], "config.PSI_BANDS")
    put("psi_hi", "PSI band: moderate below", config.PSI_BANDS[1], "config.PSI_BANDS")
    row[0] += 1
    put("sel_name", "Selected scenario", None)
    put("sel_rate", "Selected rate shock (bp)", None)
    put("sel_hpi", "Selected HPI shock (%)", None)
    put("sel_unemp", "Selected unemployment shock (pt)", None)
    row[0] += 1
    scn_hdr = row[0]
    head(ws_in, scn_hdr, ["Scenario table", "rate_bp", "hpi_pct", "unemp_pt"])
    SCN0 = scn_hdr + 1
    for i, s in enumerate(config.SHOCKS):
        for j, v in enumerate([s.name, s.rate_bp, s.hpi_pct, s.unemp_pt]):
            ws_in.cell(SCN0 + i, 1 + j, v)
    SCN1 = SCN0 + len(config.SHOCKS) - 1
    for key, j in (("sel_name", "A"), ("sel_rate", "B"), ("sel_hpi", "C"), ("sel_unemp", "D")):
        r = int(NAMES[key].split("$")[-1])
        ws_in.cell(r, 2, f"=INDEX(${j}${SCN0}:${j}${SCN1},{NAMES['sel']})")
    row[0] = SCN1 + 2
    head(ws_in, row[0], ["Prepay hinge logit term", "coef", "knot"])
    row[0] += 1
    for _, r_ in coefs.iterrows():
        t = r_[rx.TERM_COL]
        ws_in.cell(row[0], 1, t)
        ws_in.cell(row[0], 2, float(r_[rx.PP_COEF_COL]))
        if not pd.isna(r_[rx.KNOT_COL]):
            ws_in.cell(row[0], 3, float(r_[rx.KNOT_COL]))
        NAMES[f"pp_{t}"] = f"Inputs!$B${row[0]}"
        NAMES[f"ppk_{t}"] = f"Inputs!$C${row[0]}"
        row[0] += 1
    row[0] += 1
    head(ws_in, row[0], ["One-loan calculator inputs (blank = missing)", "value"])
    row[0] += 1
    for v in variables:
        ws_in.cell(row[0], 1, v)
        ws_in.cell(row[0], 2, rx.CALC_DEFAULTS[v])
        NAMES[f"calc_{v}"] = f"Inputs!$B${row[0]}"
        row[0] += 1
    ws_in.column_dimensions["A"].width = 48
    ws_in.column_dimensions["C"].width = 70
    reg("sel", "Inputs", NAMES["sel"].split("!")[1])

    # ============================================================ Scorecard (imported values + helpers)
    ws_sc["A1"] = "Scorecard bin table (inputs from Python; scorecard_table.csv)"
    ws_sc["A1"].font = BOLD
    ws_sc["A2"] = "Bins are (bin_lo, bin_hi]; the missing bin has blank bounds. Column G is the lookup helper (finite upper edges)."
    head(ws_sc, 4, ["variable", "bin_lo", "bin_hi", "woe", "coef", "points", "edge_helper"])
    head(ws_sc, 4, ["variable", "coef", "best_points", "n_bins"], col0=9)
    SC = {}
    r = 5
    for vi, v in enumerate(variables):
        g = table[table[rx.VAR_COL] == v].reset_index(drop=True)
        k = int(g[rx.LO_COL].notna().sum())
        assert k >= 2, f"{v}: scorecard lookup needs at least 2 bins"
        r0 = r
        for i, rec in g.iterrows():
            lo, hi = rec[rx.LO_COL], rec[rx.HI_COL]
            ws_sc.cell(r, 1, v)
            if i < k:
                ws_sc.cell(r, 2, "-inf" if np.isinf(lo) else float(lo))
                ws_sc.cell(r, 3, "inf" if np.isinf(hi) else float(hi))
                if i < k - 1:
                    ws_sc.cell(r, 7, f"=C{r}")
            ws_sc.cell(r, 4, float(rec[rx.WOE_COL]))
            ws_sc.cell(r, 5, float(rec[rx.COEF_COL]))
            ws_sc.cell(r, 6, float(rec[rx.POINTS_COL]))
            r += 1
        SC[v] = dict(edges=f"Scorecard!$G${r0}:$G${r0 + k - 2}", woe=f"Scorecard!$D${r0}:$D${r0 + k - 1}",
                     pts=f"Scorecard!$F${r0}:$F${r0 + k - 1}", miss_woe=f"Scorecard!$D${r0 + k}",
                     miss_pts=f"Scorecard!$F${r0 + k}", coef=f"Scorecard!$J${5 + vi}", best=f"Scorecard!$K${5 + vi}")
        ws_sc.cell(5 + vi, 9, v)
        ws_sc.cell(5 + vi, 10, f"=E{r0}")
        ws_sc.cell(5 + vi, 11, f"=MAX(F{r0}:F{r0 + k - 1})")
        ws_sc.cell(5 + vi, 12, k)
    ws_sc.column_dimensions["A"].width = 18
    ws_sc.column_dimensions["I"].width = 18

    def woe_expr(v, x):
        # bin index = count of finite edges strictly below x, plus one
        d = SC[v]
        return (f"IF(ISNUMBER({x}),INDEX({d['woe']},SUMPRODUCT(--({d['edges']}<{x}))+1),{d['miss_woe']})")

    # ============================================================ Loan_Sample
    ws_ls["A1"] = "Loan sample: 2,000 out-of-time snapshot loans (synthetic). Python columns are static; xl_ columns are live formulas."
    ws_ls["A1"].font = BOLD
    ws_ls["A2"] = "Attributes were scored in float32 in Python, so they are stored at float32 precision to land on the same bins."
    static = [LOAN_ID, SNAPSHOT_DATE] + list(FEATURES_PD) + [TARGET_PD12, CUR_UPB]
    ls_cols = {}
    heads = static + [PY_PD, PY_POINTS, PY_GBM] + [f"xl_woe_{v}" for v in variables] + \
        ["xl_margin", "xl_pd", "xl_points", "xl_rank", "xl_decile"]
    for j, h in enumerate(heads, start=1):
        ws_ls.cell(HDR, j, h).font = BOLD
        ls_cols[h] = L(j)
    cols_idx = {c: sample.columns.get_loc(c) for c in sample.columns}
    vals = sample.to_numpy(dtype=object)
    for i in range(n):
        rr = FIRST + i
        for h in static:
            v = vals[i, cols_idx[h]]
            ws_ls[f"{ls_cols[h]}{rr}"] = v.item() if hasattr(v, "item") else v
        ws_ls[f"{ls_cols[PY_PD]}{rr}"] = float(vals[i, cols_idx[SAMPLE_PY_PD]])
        ws_ls[f"{ls_cols[PY_POINTS]}{rr}"] = float(vals[i, cols_idx[SAMPLE_PY_POINTS]])
        ws_ls[f"{ls_cols[PY_GBM]}{rr}"] = float(vals[i, cols_idx[SAMPLE_GBM]])
        for v in variables:
            ws_ls[f"{ls_cols['xl_woe_' + v]}{rr}"] = "=" + woe_expr(v, f"{ls_cols[v]}{rr}")
        terms = "+".join(f"{SC[v]['coef']}*{ls_cols['xl_woe_' + v]}{rr}" for v in variables)
        m = f"{ls_cols['xl_margin']}{rr}"
        ws_ls[m] = f"={NAMES['intercept']}+{terms}"
        pdc = f"{ls_cols['xl_pd']}{rr}"
        ws_ls[pdc] = (f"=MIN(1-{NAMES['clip']},MAX({NAMES['clip']},1/(1+EXP(-({NAMES['pd_a']}*{m}+{NAMES['pd_b']})))))")
        ws_ls[f"{ls_cols['xl_points']}{rr}"] = f"={NAMES['offset']}-{NAMES['factor']}*{m}"
    pd_rng = f"${ls_cols['xl_pd']}${FIRST}:${ls_cols['xl_pd']}${last}"
    pdl = ls_cols["xl_pd"]
    for i in range(n):
        rr = FIRST + i
        ws_ls[f"{ls_cols['xl_rank']}{rr}"] = (
            f"=RANK({pdl}{rr},{pd_rng},0)+COUNTIF(${pdl}${HDR}:{pdl}{rr - 1},{pdl}{rr})")
        rk = f"{ls_cols['xl_rank']}{rr}"
        ws_ls[f"{ls_cols['xl_decile']}{rr}"] = (
            f"=IF({rk}<=Deciles!$B$5*(Deciles!$B$4+1),INT(({rk}-1)/(Deciles!$B$4+1))+1,"
            f"Deciles!$B$5+INT(({rk}-1-Deciles!$B$5*(Deciles!$B$4+1))/Deciles!$B$4)+1)")
    ws_ls.freeze_panes = f"C{FIRST}"

    def lsr(h):
        c = ls_cols[h]
        return f"Loan_Sample!${c}${FIRST}:${c}${last}"

    # ============================================================ Score_Calc
    ws_calc["A1"] = "One-loan scorecard calculator (inputs are on the Inputs sheet; blank input = missing bin)"
    ws_calc["A1"].font = BOLD
    head(ws_calc, 5, ["variable", "value", "bin_number", "woe", "coef", "coef x woe", "points", "best_bin_points",
                      "shortfall", "tiebreak_key"])
    cr0, cr1 = CALC_FIRST, CALC_FIRST + len(variables) - 1
    for i, v in enumerate(variables):
        r_ = cr0 + i
        d = SC[v]
        ws_calc.cell(r_, 1, v)
        ws_calc.cell(r_, 2, f'=IF(ISBLANK({NAMES["calc_" + v]}),"",{NAMES["calc_" + v]})')
        ws_calc.cell(r_, 3, f'=IF(ISNUMBER(B{r_}),SUMPRODUCT(--({d["edges"]}<B{r_}))+1,"missing")')
        ws_calc.cell(r_, 4, f"=IF(ISNUMBER(C{r_}),INDEX({d['woe']},C{r_}),{d['miss_woe']})")
        ws_calc.cell(r_, 5, f"={d['coef']}")
        ws_calc.cell(r_, 6, f"=D{r_}*E{r_}")
        ws_calc.cell(r_, 7, f"=IF(ISNUMBER(C{r_}),INDEX({d['pts']},C{r_}),{d['miss_pts']})")
        ws_calc.cell(r_, 8, f"={d['best']}")
        ws_calc.cell(r_, 9, f"=H{r_}-G{r_}")
        ws_calc.cell(r_, 10, f"=I{r_}-ROW()*0.000000001")
    rs = cr1 + 2
    for off, (lab, f, key) in enumerate([
            ("Sum of coef x woe", f"=SUM(F{cr0}:F{cr1})", None),
            ("Logit margin", f"={NAMES['intercept']}+B{rs}", "calc_margin"),
            ("PD (calibrated)", f"=MIN(1-{NAMES['clip']},MAX({NAMES['clip']},1/(1+EXP(-({NAMES['pd_a']}*B{rs + 1}+{NAMES['pd_b']})))))", "calc_pd"),
            ("Total score", f"={NAMES['offset']}-{NAMES['factor']}*B{rs + 1}", "calc_score"),
            ("Sum of per-variable points (equals total score)", f"=SUM(G{cr0}:G{cr1})", "calc_points_sum")]):
        ws_calc.cell(rs + off, 1, lab)
        ws_calc.cell(rs + off, 2, f)
        if key:
            reg(key, "Score_Calc", f"B{rs + off}")
    t0 = rs + 7
    ws_calc.cell(t0 - 1, 1, "Top 3 variables by points shortfall versus the best bin").font = BOLD
    head(ws_calc, t0, ["rank", "variable", "shortfall_points"])
    for k in (1, 2, 3):
        r_ = t0 + k
        ws_calc.cell(r_, 1, f"=ROW()-{t0}")
        key = f"LARGE($J${cr0}:$J${cr1},A{r_})"
        ws_calc.cell(r_, 2, f"=INDEX($A${cr0}:$A${cr1},MATCH({key},$J${cr0}:$J${cr1},0))")
        ws_calc.cell(r_, 3, f"=INDEX($I${cr0}:$I${cr1},MATCH({key},$J${cr0}:$J${cr1},0))")
        reg(f"calc_top{k}_var", "Score_Calc", f"B{r_}")
        reg(f"calc_top{k}_gap", "Score_Calc", f"C{r_}")
    ws_calc.column_dimensions["A"].width = 44

    # ============================================================ Deciles
    ws_dec["A1"] = "Decile table over the live scorecard PD (descending score, ties broken by input order, as in evaluation.decile_table)"
    ws_dec["A1"].font = BOLD
    ws_dec["A3"] = "Loans"
    ws_dec["B3"] = f"=COUNT({lsr('xl_pd')})"
    ws_dec["A4"] = "Base group size"
    ws_dec["B4"] = f"=INT(B3/{rx.N_DECILES})"
    ws_dec["A5"] = "Groups with one extra loan"
    ws_dec["B5"] = f"=MOD(B3,{rx.N_DECILES})"
    D0, D1 = DEC_HDR + 1, DEC_HDR + rx.N_DECILES
    head(ws_dec, DEC_HDR, ["decile", "n", "events", "event_rate", "mean_pd", "cum_events_share",
                           "lift", "ks_at_decile", "cum_nonevents_share", "roc_trapezoid_area",
                           "", "py_full_oot_event_rate", "py_full_oot_mean_score", "py_full_oot_ks"])
    for i in range(rx.N_DECILES):
        r_ = D0 + i
        ws_dec.cell(r_, 1, f"=ROW()-{DEC_HDR}")
        ws_dec.cell(r_, 2, f"=COUNTIFS({lsr('xl_decile')},A{r_})")
        ws_dec.cell(r_, 3, f"=SUMIFS({lsr(TARGET_PD12)},{lsr('xl_decile')},A{r_})")
        ws_dec.cell(r_, 4, f"=C{r_}/B{r_}")
        ws_dec.cell(r_, 5, f"=SUMIFS({lsr('xl_pd')},{lsr('xl_decile')},A{r_})/B{r_}")
        ws_dec.cell(r_, 6, f"=SUM(C${D0}:C{r_})/SUM(C${D0}:C${D1})")
        ws_dec.cell(r_, 7, f"=D{r_}/(SUM(C${D0}:C${D1})/SUM(B${D0}:B${D1}))")
        ws_dec.cell(r_, 8, f"=ABS(F{r_}-I{r_})")
        ws_dec.cell(r_, 9, f"=(SUM(B${D0}:B{r_})-SUM(C${D0}:C{r_}))/(SUM(B${D0}:B${D1})-SUM(C${D0}:C${D1}))")
        prev_i, prev_f = (f"I{r_ - 1}", f"F{r_ - 1}") if i else ("0", "0")
        ws_dec.cell(r_, 10, f"=(I{r_}-{prev_i})*(F{r_}+{prev_f})/2")
    # the ks column is H; keep ks_at_decile semantics aligned with the reconcile script
    ws_dec.cell(D1 + 2, 1, "Max KS (decile based)")
    ws_dec.cell(D1 + 2, 2, f"=MAX(H{D0}:H{D1})")
    reg("ks_max", "Deciles", f"B{D1 + 2}")
    ws_dec.cell(D1 + 3, 1, "AUC, approximate (decile ROC trapezoid)")
    ws_dec.cell(D1 + 3, 2, f"=SUM(J{D0}:J{D1})")
    reg("auc_approx", "Deciles", f"B{D1 + 3}")
    ws_dec.cell(D1 + 4, 1, "Gini, approximate (2 x AUC - 1)")
    ws_dec.cell(D1 + 4, 2, f"=2*B{D1 + 3}-1")
    reg("gini_approx", "Deciles", f"B{D1 + 4}")
    ws_dec.column_dimensions["A"].width = 40

    # ============================================================ Python_Ref (imported blocks)
    ws_ref["A1"] = "Python reference values (static, from the full out-of-time population unless stated)"
    ws_ref["A1"].font = BOLD
    ref_dec = pd.read_csv(config.OUT_DIR / "deciles_pd_pd_logit.csv")
    ws_ref["A3"] = "Deciles, scorecard, full OOT (deciles_pd_pd_logit.csv)"
    ws_ref["A3"].font = BOLD
    head(ws_ref, 4, list(DEC_COLS))
    for i, rec in ref_dec.iterrows():
        for j, c in enumerate(DEC_COLS, start=1):
            ws_ref.cell(5 + i, j, float(rec[c]))
    PR_DEC = 5
    ws_ref["A17"] = "PSI expected bins (psi_train_bins.csv); -inf shown as -1 and inf as 2 (PD domain bounds)"
    ws_ref["A17"].font = BOLD
    head(ws_ref, 18, ["bin", "lo", "hi", "train_share"])
    for i, rec in psi_bins.reset_index(drop=True).iterrows():
        lo, hi = rec[rx.PSI_LO_COL], rec[rx.PSI_HI_COL]
        ws_ref.cell(19 + i, 1, int(rec["bin"]))
        ws_ref.cell(19 + i, 2, -1.0 if np.isinf(lo) else float(lo))
        ws_ref.cell(19 + i, 3, 2.0 if np.isinf(hi) else float(hi))
        ws_ref.cell(19 + i, 4, float(rec[rx.PSI_SHARE_COL]))
    PR_PSI = 19
    csi = pd.read_csv(config.OUT_DIR / "psi_csi.csv")
    ws_ref["A31"] = "Python PSI and CSI, train vs full OOT (psi_csi.csv)"
    ws_ref["A31"].font = BOLD
    head(ws_ref, 32, list(csi.columns))
    for i, rec in csi.iterrows():
        for j, c in enumerate(csi.columns, start=1):
            v = rec[c]
            ws_ref.cell(33 + i, j, v if isinstance(v, str) else float(v))
    r_s = 33 + len(csi) + 2
    s12 = pd.read_csv(config.OUT_DIR / "scenario_12m.csv")
    ws_ref.cell(r_s, 1, "Python GBM scenario summary (scenario_12m.csv; PD and prepay GBMs, 60-month projection)").font = BOLD
    head(ws_ref, r_s + 1, list(s12.columns))
    for i, rec in s12.iterrows():
        for j, c in enumerate(s12.columns, start=1):
            v = rec[c]
            ws_ref.cell(r_s + 2 + i, j, v if isinstance(v, str) else float(v))
    S12_0, S12_1 = r_s + 2, r_s + 1 + len(s12)
    bm = pd.read_csv(config.OUT_DIR / "benchmark_pd.csv")
    r_b = S12_1 + 3
    ws_ref.cell(r_b, 1, "PD model benchmark (benchmark_pd.csv)").font = BOLD
    head(ws_ref, r_b + 1, list(bm.columns))
    for i, rec in bm.iterrows():
        for j, c in enumerate(bm.columns, start=1):
            v = rec[c]
            ws_ref.cell(r_b + 2 + i, j, v if isinstance(v, str) else (None if pd.isna(v) else float(v)))
    ws_ref.column_dimensions["A"].width = 30

    # Deciles side-by-side (formulas pulling the static reference)
    for i in range(rx.N_DECILES):
        r_ = D0 + i
        ws_dec.cell(r_, 12, f"=Python_Ref!D{PR_DEC + i}")
        ws_dec.cell(r_, 13, f"=Python_Ref!E{PR_DEC + i}")
        ws_dec.cell(r_, 14, f"=Python_Ref!H{PR_DEC + i}")

    # ============================================================ PSI
    ws_psi["A1"] = "Population stability index of the scorecard PD: training quantile bins vs the 2,000-loan sample (live)"
    ws_psi["A1"].font = BOLD
    ws_psi["A3"] = "Share floor"
    ws_psi["B3"] = f"={NAMES['psi_floor']}"
    ws_psi["A4"] = "Band thresholds (stable below / moderate below)"
    ws_psi["B4"] = f"={NAMES['psi_lo']}"
    ws_psi["C4"] = f"={NAMES['psi_hi']}"
    head(ws_psi, 6, ["bin", "lo (exclusive)", "hi (inclusive)", "expected_share", "actual_share", "actual_count",
                     "expected_floored", "actual_floored", "psi_contribution"])
    P0, P1 = 7, 7 + rx.N_DECILES - 1
    for i in range(rx.N_DECILES):
        r_ = P0 + i
        ws_psi.cell(r_, 1, f"=Python_Ref!A{PR_PSI + i}")
        ws_psi.cell(r_, 2, f"=Python_Ref!B{PR_PSI + i}")
        ws_psi.cell(r_, 3, f"=Python_Ref!C{PR_PSI + i}")
        ws_psi.cell(r_, 4, f"=Python_Ref!D{PR_PSI + i}")
        ws_psi.cell(r_, 6, f"=SUMPRODUCT(({lsr('xl_pd')}>B{r_})*({lsr('xl_pd')}<=C{r_}))")
        ws_psi.cell(r_, 5, f"=F{r_}/SUM(F${P0}:F${P1})")
        ws_psi.cell(r_, 7, f"=MAX(D{r_},$B$3)")
        ws_psi.cell(r_, 8, f"=MAX(E{r_},$B$3)")
        ws_psi.cell(r_, 9, f"=(H{r_}-G{r_})*LN(H{r_}/G{r_})")
    ws_psi.cell(P1 + 2, 1, "Total PSI")
    ws_psi.cell(P1 + 2, 2, f"=SUM(I{P0}:I{P1})")
    reg("psi_total", "PSI", f"B{P1 + 2}")
    ws_psi.cell(P1 + 3, 1, "Python PSI, train vs full OOT (pd_logit)")
    ws_psi.cell(P1 + 3, 2, f'=INDEX(Python_Ref!$C$33:$C${32 + len(csi)},MATCH("pd_logit:train_vs_oot",Python_Ref!$B$33:$B${32 + len(csi)},0))')
    ws_psi.cell(P1 + 4, 1, "Note: the training distribution is stored only for the score, so per-variable PSI is not recomputed live; Python CSI is on Python_Ref.")
    ws_psi.column_dimensions["A"].width = 40

    # ============================================================ Calibration
    ws_cal["A1"] = "Calibration of the live scorecard PD by decile (decile 1 = highest predicted PD)"
    ws_cal["A1"].font = BOLD
    head(ws_cal, 3, ["decile", "n", "mean_predicted_pd", "observed_rate", "predicted_minus_observed",
                     "expected_events", "events", "hosmer_lemeshow_term"])
    K0, K1 = 4, 4 + rx.N_DECILES - 1
    for i in range(rx.N_DECILES):
        r_, dr = K0 + i, D0 + i
        ws_cal.cell(r_, 1, f"=Deciles!A{dr}")
        ws_cal.cell(r_, 2, f"=Deciles!B{dr}")
        ws_cal.cell(r_, 3, f"=Deciles!E{dr}")
        ws_cal.cell(r_, 4, f"=Deciles!D{dr}")
        ws_cal.cell(r_, 5, f"=C{r_}-D{r_}")
        ws_cal.cell(r_, 6, f"=B{r_}*C{r_}")
        ws_cal.cell(r_, 7, f"=Deciles!C{dr}")
        ws_cal.cell(r_, 8, f"=(G{r_}-F{r_})^2/(B{r_}*C{r_}*(1-C{r_}))")
    labs = [("Brier score (live)", f"=SUMPRODUCT(({lsr('xl_pd')}-{lsr(TARGET_PD12)})^2)/COUNT({lsr('xl_pd')})", "brier"),
            ("Brier of a constant base-rate forecast", f"=(SUM(G{K0}:G{K1})/SUM(B{K0}:B{K1}))*(1-SUM(G{K0}:G{K1})/SUM(B{K0}:B{K1}))", "brier_base"),
            ("Hosmer-Lemeshow statistic (10 groups)", f"=SUM(H{K0}:H{K1})", "hl"),
            ("Total expected events", f"=SUM(F{K0}:F{K1})", "exp_events"),
            ("Total observed events", f"=SUM(G{K0}:G{K1})", "obs_events"),
            ("Expected / observed", f"=B{K1 + 5}/B{K1 + 6}", "ratio")]
    for i, (lab, f, key) in enumerate(labs):
        ws_cal.cell(K1 + 2 + i, 1, lab)
        ws_cal.cell(K1 + 2 + i, 2, f)
        reg(key, "Calibration", f"B{K1 + 2 + i}")
    ws_cal.column_dimensions["A"].width = 42

    # ============================================================ Curves_Python
    curves = pd.read_csv(config.OUT_DIR / "scenario_curves.csv")
    ws_cur["A1"] = "Python GBM scenario curves (scenario_curves.csv; static reference, 60 months per scenario)"
    ws_cur["A1"].font = BOLD
    head(ws_cur, 3, list(curves.columns))
    for i, rec in enumerate(curves.itertuples(index=False)):
        for j, v in enumerate(rec, start=1):
            ws_cur.cell(4 + i, j, v if isinstance(v, str) else float(v))

    # ============================================================ Scenario
    ws_scn["A1"] = "Scenario: selected shock applied to the 2,000-loan sample (scorecard re-binned on shocked attributes; prepay hinge logit)"
    ws_scn["A1"].font = BOLD
    ws_scn["A2"] = ("Shock mapping: unemp and unemp_chg_12m + unemp_pt; mtm_ltv / (1 + hpi_pct/100); incentive - rate_bp/100. "
                    "Full shock applied at the snapshot (no ramp). The scorecard uses these four macro-sensitive variables directly, "
                    "so no separate uplift table is needed.")
    ws_scn["A3"] = "Selected scenario"
    ws_scn["B3"] = f"={NAMES['sel_name']}"
    head(ws_scn, 5, ["Metric", "Base (no shock)", "Selected scenario", "Change"])
    S0 = SCEN_HDR + 1
    S1 = S0 + n - 1
    sc_cols = {}
    scen_heads = [LOAN_ID, CUR_UPB, "unemp_s", "unemp_chg_s", "mtm_ltv_s", "incentive_s"] + \
        [f"woe_s_{v}" for v in SCEN_SHOCK_VARS if v in variables] + \
        ["margin_s", "pd_s", "el_base", "el_s", "pp_margin_base", "smm_base", "pp_margin_s", "smm_s"]
    for j, h in enumerate(scen_heads, start=1):
        ws_scn.cell(SCEN_HDR, j, h).font = BOLD
        sc_cols[h] = L(j)

    def sr(h):
        return f"${sc_cols[h]}${S0}:${sc_cols[h]}${S1}"

    def pp_margin(inc, mtm, rr, ls_r):
        a = lambda k: NAMES[f"pp_{k}"]
        hinge = "+".join(f"{a(t)}*MAX({inc}-{NAMES['ppk_' + t]},0)" for t in coefs[rx.TERM_COL] if str(t).startswith("hinge_"))
        return (f"{a('const')}+{hinge}+{a('burnout')}*{NAMES['burnout']}"
                f"+{a('seasoning_ramp')}*MIN(Loan_Sample!{ls_cols[AGE]}{ls_r}/{NAMES['season']},1)"
                f"+{a(MTM_LTV)}*MIN({mtm},{NAMES['mtm_cap']})+{a('fico')}*Loan_Sample!{ls_cols['fico']}{ls_r}")

    for i in range(n):
        r_, lr = S0 + i, FIRST + i
        c = sc_cols
        ws_scn[f"{c[LOAN_ID]}{r_}"] = f"=Loan_Sample!{ls_cols[LOAN_ID]}{lr}"
        ws_scn[f"{c[CUR_UPB]}{r_}"] = f"=Loan_Sample!{ls_cols[CUR_UPB]}{lr}"
        ws_scn[f"{c['unemp_s']}{r_}"] = f"=Loan_Sample!{ls_cols[UNEMP]}{lr}+{NAMES['sel_unemp']}"
        ws_scn[f"{c['unemp_chg_s']}{r_}"] = f"=Loan_Sample!{ls_cols[UNEMP_CHG_12M]}{lr}+{NAMES['sel_unemp']}"
        ws_scn[f"{c['mtm_ltv_s']}{r_}"] = f"=Loan_Sample!{ls_cols[MTM_LTV]}{lr}/(1+{NAMES['sel_hpi']}/100)"
        ws_scn[f"{c['incentive_s']}{r_}"] = f"=Loan_Sample!{ls_cols[INCENTIVE]}{lr}-{NAMES['sel_rate']}/100"
        xmap = {UNEMP: "unemp_s", UNEMP_CHG_12M: "unemp_chg_s", MTM_LTV: "mtm_ltv_s", INCENTIVE: "incentive_s"}
        terms = []
        for v in variables:
            if v in xmap:
                ws_scn[f"{c['woe_s_' + v]}{r_}"] = "=" + woe_expr(v, f"{c[xmap[v]]}{r_}")
                terms.append(f"{SC[v]['coef']}*{c['woe_s_' + v]}{r_}")
            else:
                terms.append(f"{SC[v]['coef']}*Loan_Sample!{ls_cols['xl_woe_' + v]}{lr}")
        ws_scn[f"{c['margin_s']}{r_}"] = f"={NAMES['intercept']}+" + "+".join(terms)
        ws_scn[f"{c['pd_s']}{r_}"] = (f"=MIN(1-{NAMES['clip']},MAX({NAMES['clip']},"
                                      f"1/(1+EXP(-({NAMES['pd_a']}*{c['margin_s']}{r_}+{NAMES['pd_b']})))))")
        ws_scn[f"{c['el_base']}{r_}"] = f"=Loan_Sample!{pdl}{lr}*{NAMES['lgd']}*{c[CUR_UPB]}{r_}"
        ws_scn[f"{c['el_s']}{r_}"] = f"={c['pd_s']}{r_}*{NAMES['lgd']}*{c[CUR_UPB]}{r_}"
        ws_scn[f"{c['pp_margin_base']}{r_}"] = "=" + pp_margin(f"Loan_Sample!{ls_cols[INCENTIVE]}{lr}",
                                                                f"Loan_Sample!{ls_cols[MTM_LTV]}{lr}", r_, lr)
        ws_scn[f"{c['smm_base']}{r_}"] = (f"=MIN(1-{NAMES['clip']},MAX({NAMES['clip']},"
                                          f"1/(1+EXP(-({NAMES['pp_a']}*{c['pp_margin_base']}{r_}+{NAMES['pp_b']})))))")
        ws_scn[f"{c['pp_margin_s']}{r_}"] = "=" + pp_margin(f"{c['incentive_s']}{r_}", f"{c['mtm_ltv_s']}{r_}", r_, lr)
        ws_scn[f"{c['smm_s']}{r_}"] = (f"=MIN(1-{NAMES['clip']},MAX({NAMES['clip']},"
                                       f"1/(1+EXP(-({NAMES['pp_a']}*{c['pp_margin_s']}{r_}+{NAMES['pp_b']})))))")
    upb = sr(CUR_UPB)
    pdb = lsr("xl_pd")
    summ = [
        ("Loans", f"=COUNT({pdb})", f"=COUNT({sr('pd_s')})", None),
        ("Total UPB (USD)", f"=SUM({upb})", f"=SUM({upb})", None),
        ("Expected 12-month defaults (sum of PD)", f"=SUM({pdb})", f"=SUM({sr('pd_s')})", "exp_defaults"),
        ("UPB-weighted 12-month PD", f"=SUMPRODUCT({pdb},{upb})/SUM({upb})", f"=SUMPRODUCT({sr('pd_s')},{upb})/SUM({upb})", "pd_w"),
        ("Expected loss = PD x LGD x UPB (USD)", f"=SUM({sr('el_base')})", f"=SUM({sr('el_s')})", "el_usd"),
        ("Expected loss, % of UPB", f"=SUM({sr('el_base')})/SUM({upb})", f"=SUM({sr('el_s')})/SUM({upb})", None),
        ("UPB-weighted SMM (monthly prepay rate)", f"=SUMPRODUCT({sr('smm_base')},{upb})/SUM({upb})",
         f"=SUMPRODUCT({sr('smm_s')},{upb})/SUM({upb})", "smm_w"),
        ("12-month prepay at constant SMM = 1-(1-SMM)^12 (CPR)", "=1-(1-B12)^12", "=1-(1-C12)^12", "prepay_12m"),
        ("12-month default rate (CDR proxy) = UPB-weighted PD", "=B9", "=C9", None),
    ]
    for i, (lab, fb, fs, key) in enumerate(summ):
        r_ = 6 + i
        ws_scn.cell(r_, 1, lab)
        ws_scn.cell(r_, 2, fb)
        ws_scn.cell(r_, 3, fs)
        ws_scn.cell(r_, 4, f"=C{r_}-B{r_}")
        if key:
            reg(f"base_{key}", "Scenario", f"B{r_}")
            reg(f"sel_{key}", "Scenario", f"C{r_}")
    ws_scn.cell(16, 1, "Python GBM reference for the selected scenario (static curves, 12 months; context only)").font = BOLD
    for i, c_ in enumerate(SCEN12_COLS):
        j = list(s12.columns).index(c_)
        colL = L(j + 1)
        ws_scn.cell(17 + i, 1, c_)
        ws_scn.cell(17 + i, 3, f"=INDEX(Python_Ref!${colL}${S12_0}:${colL}${S12_1},MATCH($B$3,Python_Ref!$A${S12_0}:$A${S12_1},0))")
    ws_scn.cell(21, 1, "Note: the GBM PD is insensitive to the rate shock; the scorecard moves only through its incentive variable.")
    ws_scn.column_dimensions["A"].width = 56

    # ============================================================ Checks
    ws_chk["A1"] = "Integrity checks (1 = pass)"
    ws_chk["A1"].font = BOLD
    head(ws_chk, 3, ["check", "value", "pass"])
    dn = f"Deciles!$B${D0}:$B${D1}"
    checks = [
        ("Loan count equals decile count total", f"=Deciles!B3-SUM({dn})", "=IF(ABS(B{r})<0.5,1,0)"),
        ("Decile counts sum to the sample size", f"=SUM({dn})", f"=IF(B{{r}}=COUNTA({lsr(LOAN_ID)}),1,0)"),
        ("Events across deciles equal sample events", f"=SUM(Deciles!$C${D0}:$C${D1})-SUM({lsr(TARGET_PD12)})", "=IF(ABS(B{r})<0.5,1,0)"),
        ("All base PD within [0,1]", f"=MIN({pdb})", f"=IF(AND(B{{r}}>=0,MAX({pdb})<=1),1,0)"),
        ("All stressed PD within [0,1]", f"=MIN(Scenario!{sr('pd_s')})", f"=IF(AND(B{{r}}>=0,MAX(Scenario!{sr('pd_s')})<=1),1,0)"),
        ("All SMM within [0,1]", f"=MIN(Scenario!{sr('smm_s')})", f"=IF(AND(B{{r}}>=0,MAX(Scenario!{sr('smm_s')},Scenario!{sr('smm_base')})<=1),1,0)"),
        ("Ranks are a permutation of 1..N", f"=SUM({lsr('xl_rank')})-Deciles!B3*(Deciles!B3+1)/2", "=IF(ABS(B{r})<0.5,1,0)"),
        ("Expected PSI shares sum to 1", f"=SUM(PSI!$D${P0}:$D${P1})-1", "=IF(ABS(B{r})<0.000001,1,0)"),
        ("Actual PSI shares sum to 1", f"=SUM(PSI!$E${P0}:$E${P1})-1", "=IF(ABS(B{r})<0.000000001,1,0)"),
        ("Cumulative event capture ends at 1", f"=Deciles!F{D1}-1", "=IF(ABS(B{r})<0.000000001,1,0)"),
        ("Expected events equal sum of PD", f"=Calibration!B{K1 + 5}-SUM({pdb})", "=IF(ABS(B{r})<0.000001,1,0)"),
        ("Live PD matches the exported Python PD (max abs diff)",
         f"=SUMPRODUCT(MAX(ABS({lsr('xl_pd')}-{lsr(PY_PD)})))", "=IF(B{r}<0.000001,1,0)"),
        ("Live points match the exported Python points (max abs diff)",
         f"=SUMPRODUCT(MAX(ABS({lsr('xl_points')}-{lsr(PY_POINTS)})))", "=IF(B{r}<0.001,1,0)"),
        ("Calculator per-variable points sum to total score", f"=Score_Calc!B{rs + 4}-Score_Calc!B{rs + 3}", "=IF(ABS(B{r})<0.000001,1,0)"),
    ]
    for i, (lab, fv, fp) in enumerate(checks):
        r_ = 4 + i
        ws_chk.cell(r_, 1, lab)
        ws_chk.cell(r_, 2, fv)
        ws_chk.cell(r_, 3, fp.format(r=r_))
    c1 = 4 + len(checks) - 1
    ws_chk.cell(c1 + 2, 1, "All checks pass (1 = yes)").font = BOLD
    ws_chk.cell(c1 + 2, 2, f"=IF(SUM(C4:C{c1})=COUNT(C4:C{c1}),1,0)")
    reg("checks_all", "Checks", f"B{c1 + 2}")
    ws_chk.column_dimensions["A"].width = 62

    # ============================================================ Conclusions (static text)
    dec, psi, cal, base = ref["dec"], ref["psi"], ref["cal"], ref["base"]
    band = "stable" if psi["total"] < config.PSI_BANDS[0] else ("moderate" if psi["total"] < config.PSI_BANDS[1] else "significant")
    ratio = cal["exp_events"] / cal["events"]
    cal_word = ("well calibrated in aggregate" if abs(ratio - 1) < 0.10 else
                ("moderately miscalibrated in aggregate" if abs(ratio - 1) < 0.25 else "poorly calibrated in aggregate"))
    lines = [
        "Conclusions (static text written by build_workbook.py from the computed numbers; rerun the build to refresh)",
        "",
        "All data are synthetic, simulated from a known generating process. These results show the method works, not how real loans behave.",
        "",
        f"Rank ordering: on {len(sample):,} out-of-time loans the decile KS is {dec['ks_max']:.3f}; the decile-trapezoid Gini is {dec['gini_approx']:.3f} "
        f"(approximate, exact AUC is {rx_auc(ref):.3f}). The top decile holds {dec['cum1'][0]:.0%} of defaults, a lift of {dec['lift'][0]:.1f}x.",
        f"Stability: PSI of the scorecard PD against the training bins is {psi['total']:.4f}, in the {band} band "
        f"(thresholds {config.PSI_BANDS[0]} and {config.PSI_BANDS[1]}). Per-variable CSI is on Python_Ref, not recomputed here.",
        f"Calibration: mean predicted PD is {base['pd'].mean():.3%} against an observed rate of {cal['events'] / len(sample):.3%} "
        f"(expected/observed events {ratio:.2f}); the model is {cal_word}. Brier {cal['brier']:.5f} versus {cal['brier_base']:.5f} for a constant base-rate forecast. "
        f"Hosmer-Lemeshow statistic {cal['hl']:.1f} on 10 groups; with only {int(cal['events'])} events the deciles are noisy.",
        "",
        "Scenario results on the sample (scorecard re-binned on shocked attributes; prepay hinge logit with BURNOUT fixed at "
        f"{p['burnout']:g} months, solved to match the Python GBM base 12-month prepay; an approximation):",
    ]
    for s, res in zip(config.SHOCKS, ref["scen"]):
        lines.append(f"  {s.name}: expected 12-month defaults {res['exp_defaults']:.1f}, UPB-weighted PD {res['pd_w']:.3%}, "
                     f"expected loss ${res['el_usd']:,.0f} ({res['el_pct']:.3%} of UPB), SMM {res['smm_w']:.4%}, 12-month prepay {res['prepay_12m']:.2%}.")
    lines += ["",
              f"The adverse combined case raises expected defaults from {ref['scen'][0]['exp_defaults']:.1f} to {ref['scen'][5]['exp_defaults']:.1f} "
              f"({ref['scen'][5]['exp_defaults'] / ref['scen'][0]['exp_defaults'] - 1:+.0%}); unemployment and house-price shocks drive scorecard PD, rate shocks only slightly.",
              "Rate shocks move prepay a great deal in this hinge logit (a 200bp rise cuts incentive and the 12-month prepay rate falls sharply), "
              "whereas the Python GBM curves barely move. Treat the rate-scenario prepay figures as a logit extrapolation, not a forecast.",
              "Limits: one scenario step is applied at the snapshot without a ramp, LGD is a flat assumption, and the prepay side is a single-month hazard "
              "annualised at a constant SMM, not the 60-month Python projection (see Curves_Python).",
              "The Python comparison is in Checks (live formulas versus exported Python scores) and in outputs/excel_reconciliation.json."]
    for i, t in enumerate(lines, start=1):
        ws_con.cell(i, 1, t)
    ws_con["A1"].font = BOLD
    ws_con.column_dimensions["A"].width = 160

    # ============================================================ README
    readme = [
        "Mortgage delinquency and prepayment model: Excel companion (synthetic data)",
        "",
        "Purpose: a live, auditable version of the scorecard, decile, PSI, calibration and scenario calculations for 2,000 out-of-time loans.",
        "All loans are simulated. Nothing here describes real borrowers.",
        "",
        "Static values from Python: Scorecard bins, Loan_Sample attributes and Python_Ref, Curves_Python, and the model scalars on Inputs.",
        "Live formulas: Loan_Sample xl_ columns, Score_Calc, Deciles, PSI, Calibration, Scenario and Checks.",
        "Conclusions is static text written at build time.",
        "",
        "How to use: change the scenario selector on Inputs (1 to 6) to re-run the shocked PD, expected loss and prepay summary on Scenario.",
        "Change the one-loan calculator inputs on Inputs to score a single loan on Score_Calc (clear a cell to test the missing bin).",
        "Checks cell 'All checks pass' must read 1.",
        "",
        "Approximations: the prepay side holds BURNOUT fixed at a value solved to match the Python GBM base prepay (no history in the sample); decile AUC and Gini are trapezoid approximations; "
        "the scenario applies the full shock at the snapshot with no ramp.",
        "Python reconciliation is run by mortgage_risk.reconcile_excel against a LibreOffice recalculation; see outputs/excel_reconciliation.json.",
    ]
    for i, t in enumerate(readme, start=1):
        ws_readme.cell(i, 1, t)
    ws_readme["A1"].font = BOLD
    ws_readme.column_dimensions["A"].width = 140

    n_formulas = sum(1 for ws in wb for r_ in ws.iter_rows() for c in r_ if isinstance(c.value, str) and c.value.startswith("="))
    meta = dict(addr=addr, first=FIRST, last=last, ls_cols=ls_cols, deciles=dict(first=D0), psi=dict(first=P0),
                scen_first=S0, scen_last=S1, scen_cols=sc_cols, n_formulas=n_formulas, names=NAMES)
    wb.save(out_file)
    if write_meta:
        OUT_FILE.with_suffix(".meta.json").write_text(json.dumps(meta, indent=1))
    return meta


def rx_auc(ref):
    return ref["auc_exact"]


if __name__ == "__main__":
    m = build()
    print("saved", OUT_FILE, "formulas:", m["n_formulas"])
