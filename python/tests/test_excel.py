"""Workbook tests: plain formatting, author, formula/typed-value separation, static conclusions, LibreOffice reconciliation."""
import json
import zipfile

import openpyxl
import pytest

from mortgage_risk import reconcile_excel as rx

# ===== CONFIG (user inputs) =====
FORMULA_SHEETS = ("Score_Calc", "Deciles", "PSI", "Calibration", "Scenario", "Checks")
STATIC_SHEETS = ("README", "Conclusions")
LIVE_PREFIX = "xl_"
MIN_CASES = 2
# ===== END CONFIG =====

needs_lo = pytest.mark.skipif(rx.find_soffice() is None, reason="LibreOffice not installed")


@pytest.fixture(scope="module")
def report():
    return rx.build_and_reconcile(quick=True)


@pytest.fixture(scope="module")
def wb(report):
    return openpyxl.load_workbook(rx.XLSX)


@needs_lo
def test_reconciliation_passes(report):
    assert report["all_passed"], [c["issues"] for c in report["cases"] if not c["passed"]]
    assert report["n_cases"] >= MIN_CASES
    assert json.loads(rx.REPORT.read_text())["all_passed"]


@needs_lo
def test_plain_formatting(wb):
    for ws in wb:
        assert ws.sheet_properties.tabColor is None, ws.title
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                assert c.fill is None or c.fill.fill_type in (None, "none"), (ws.title, c.coordinate)
                color = c.font.color
                assert color is None or color.type != "rgb" or color.rgb in ("FF000000", "00000000"), (ws.title, c.coordinate)


@needs_lo
def test_author_metadata(wb):
    assert wb.properties.creator == "Deepak Chaudhary"
    assert wb.properties.lastModifiedBy == "Deepak Chaudhary"
    with zipfile.ZipFile(rx.XLSX) as z:
        core = z.read("docProps/core.xml").decode("utf-8")
    assert "Deepak Chaudhary" in core


@needs_lo
def test_formula_sheets_have_no_typed_numbers(wb):
    for name in FORMULA_SHEETS:
        for row in wb[name].iter_rows():
            for c in row:
                v = c.value
                is_number = isinstance(v, (int, float)) and not isinstance(v, bool)
                assert not is_number, (name, c.coordinate, v)


@needs_lo
def test_loan_sample_live_columns_are_formulas(wb):
    ws = wb["Loan_Sample"]
    meta = json.loads(rx.META.read_text())
    for h, letter in meta["ls_cols"].items():
        if not h.startswith(LIVE_PREFIX):
            continue
        for r in (meta["first"], (meta["first"] + meta["last"]) // 2, meta["last"]):
            v = ws[f"{letter}{r}"].value
            assert isinstance(v, str) and v.startswith("="), (h, r, v)


@needs_lo
def test_conclusions_static_no_if_branching(wb):
    for name in STATIC_SHEETS:
        for row in wb[name].iter_rows():
            for c in row:
                assert not (isinstance(c.value, str) and c.value.startswith("=")), (name, c.coordinate)
    # no formula anywhere uses IF to build a conclusion string on the Conclusions sheet (checked above); sheet exists
    assert "Conclusions" in wb.sheetnames


@needs_lo
def test_no_em_dashes_in_workbook(wb):
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    assert chr(0x2014) not in c.value and chr(0x2013) not in c.value, (ws.title, c.coordinate)


@needs_lo
def test_checks_cell_reports_pass(report):
    wbv = openpyxl.load_workbook(rx.XLSX, data_only=True)
    meta = json.loads(rx.META.read_text())
    sheet, ref = meta["addr"]["checks_all"].split("!")
    assert wbv[sheet][ref].value == 1
