"""Guards on the generated documentation (docx, markdown, README); reads files only."""
import re
import sys
import zipfile
from pathlib import Path

import pytest

# ===== CONFIG (user inputs) =====
ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
DOCX, MD, README = DOCS / "Mortgage_Risk_Model_Documentation.docx", DOCS / "Mortgage_Risk_Model_Documentation.md", ROOT / "README.md"
AUTHOR = "Deepak Chaudhary"
EM_DASH = chr(0x2014)
DASH_SUBSTITUTE = re.compile(r"(?<=\w)--(?=\w)|\s--\s|" + EM_DASH)
CITE_TOKENS = ("Siddiqi", "Chen and Guestrin", "Ke et al", "Lundberg and Lee", "DeLong, DeLong", "Federal Reserve and OCC")
# ===== END CONFIG =====

pytestmark = pytest.mark.skipif(not DOCX.exists(), reason="docs not built")
sys.path.insert(0, str(DOCS))


def _docx_text():
    from docx import Document
    d = Document(str(DOCX))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [c.text for r in t.rows for c in r.cells]
    return "\n".join(parts)


def test_files_exist():
    assert DOCX.exists() and MD.exists() and README.exists()


def test_no_em_or_double_hyphen_dashes():
    for name, text in (("docx", _docx_text()), ("md", MD.read_text(encoding="utf-8")), ("readme", README.read_text(encoding="utf-8"))):
        assert not DASH_SUBSTITUTE.search(text), f"dash substitute in {name}: {DASH_SUBSTITUTE.search(text)}"
        assert "\ufffd" not in text


def test_heading_styles_have_no_color():
    xml = zipfile.ZipFile(DOCX).read("word/styles.xml").decode("utf-8")
    for block in re.findall(r"<w:style [^>]*w:styleId=\"Heading\d\".*?</w:style>", xml, flags=re.S):
        assert "<w:color" not in block and "themeColor" not in block


def test_author_properties():
    from docx import Document
    cp = Document(str(DOCX)).core_properties
    assert cp.author == AUTHOR and cp.last_modified_by == AUTHOR


def test_counts_match_markdown():
    import build_docs as bd
    assert bd.md_stats(MD.read_text(encoding="utf-8")) == bd.docx_stats(DOCX)


def test_image_paths_exist():
    text = MD.read_text(encoding="utf-8")
    paths = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", text)
    assert paths
    for p in paths:
        assert (DOCS / p).resolve().exists(), p


def test_at_most_one_citation_per_paragraph():
    for para in re.split(r"\n\s*\n", MD.read_text(encoding="utf-8")):
        n = sum(para.count(t) for t in CITE_TOKENS)
        assert n <= 1, para[:120]


def test_no_unresolved_placeholders():
    assert "@@" not in MD.read_text(encoding="utf-8") and "@@" not in README.read_text(encoding="utf-8")
